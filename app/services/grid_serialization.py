"""
Motor de serialização da planilha interativa (spec 009).

Converte entre três representações:
  (a) XLSX (bytes)      — o que a Senior/cliente exporta e o sistema gera.
  (b) grid (dict JSON)  — modelo neutro consumido pelo Univer no front.
  (c) BillingModel /    — o modelo canônico do sistema (estrutura C1,
      estrutura C1           colunas, campos_config, percentuais).

Reutiliza `excel_export.py` e `model_structure.py` por IMPORT (sem editá-los):
  - `billing_to_femsa_excel` / `_render_por_estrutura` produzem os bytes .xlsx
    fiéis do faturamento; o preview do grid é SEMPRE derivado desses bytes
    (`xlsx_to_grid`), garantindo "o que você vê é o que exporta" (SC-0).
  - `derive_colunas` / `validate_estrutura` / `GERAL_COLUMNS` guiam a validação
    das operações (contrato grid-ops v1).

Formato do "grid" (contrato de serialização desta spec — neutro p/ Univer):

{
  "sheet": "Faturamento Julho",
  "n_rows": 5,
  "n_cols": 78,
  "cells": [                        # matriz densa row-major; None = célula vazia
    [{"v": "Empresa"}, {"v": "Mês Referência"}, ...],
    ...
  ],
  "merges": ["A1:B1", ...],         # ranges A1 mesclados (do cabeçalho)
  "meta": {                         # metadados de estrutura (SEM PII)
    "header_rows": [1],
    "data_row": 2,
    "colunas": [ {header, tipo, fonte/template/valor}, ... ]
  }
}

Cada célula é `{"v": <escalar JSON>}` ou `None`. Fórmulas ficam como string
começando com "=". Sem estilo aqui (o Univer aplica tema próprio; o estilo real
mora no .xlsx e é preservado pelos renderizadores de excel_export).

Regra de PII (§5 do contrato): `get_snapshot` NUNCA devolve dados de linha —
apenas metadados de estrutura (colunas, tipos, fórmulas, parâmetros) e, no
máximo, uma linha de exemplo SINTÉTICA. `model_to_grid` também é sem PII (só
cabeçalhos do modelo). Já o preview/export (`xlsx_to_grid` sobre bytes reais)
trafega dados reais no BACKEND, nunca no snapshot do prompt.
"""
from __future__ import annotations

import copy
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string

# Reuso por import (NÃO editar estes módulos):
from app.services.excel_export import (
    GERAL_COLUMNS,
    FEMSA_COLUMNS,
    _render_por_estrutura,
)
from app.services.model_structure import (
    derive_colunas,
    validate_estrutura,
    COLUNAS_CONHECIDAS,
)

# Colunas canônicas conhecidas do sistema (para validar fontes de add_column).
_KNOWN = set(COLUNAS_CONHECIDAS)

# Parâmetros aceitos por set_model_param e onde gravam no BillingModel dict.
_PARAM_TO_FIELD = {
    "encargos_pct": "encargos_pct",
    "taxa_adm_pct": "taxa_adm_pct",
    "imposto_pct": "imposto_pct",
    "salario_formula": "salario_formula",
}
_PCT_PARAMS = {"encargos_pct", "taxa_adm_pct", "imposto_pct"}


# ---------------------------------------------------------------------------
# ApplyResult (contrato §6)
# ---------------------------------------------------------------------------
@dataclass
class ApplyResult:
    ok: bool = True
    applied: int = 0
    warnings: List[str] = field(default_factory=list)
    errors: List[str] = field(default_factory=list)
    model: Optional[Dict[str, Any]] = None

    def to_dict(self) -> Dict[str, Any]:
        return {
            "ok": self.ok,
            "applied": self.applied,
            "warnings": list(self.warnings),
            "errors": list(self.errors),
            "model": self.model,
        }


# ---------------------------------------------------------------------------
# Helpers de valor
# ---------------------------------------------------------------------------
def _cell_value(v: Any) -> Any:
    """Normaliza o valor de uma célula openpyxl para um escalar JSON."""
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, (int, float, bool, str)):
        return v
    return str(v)


def _norm_num(v: Any) -> Any:
    """Trata NaN/inf e ints-como-float para comparação tolerante no round-trip."""
    if isinstance(v, bool):
        return v
    if isinstance(v, float):
        if v != v:  # NaN
            return None
        # 5.0 == 5 na igualdade tolerante
        if v.is_integer():
            return int(v)
        return round(v, 6)
    if isinstance(v, int):
        return v
    if isinstance(v, str):
        return v.strip()
    return v


# ---------------------------------------------------------------------------
# XLSX -> grid
# ---------------------------------------------------------------------------
def xlsx_to_grid(conteudo: bytes, sheet: Optional[str] = None) -> Dict[str, Any]:
    """Converte os bytes de um .xlsx no dict "grid" neutro (para o Univer).

    Lê a aba principal (primeira, ou `sheet` se informado), produz uma matriz
    densa de células `{"v": ...}`/None do canto A1 até (max_row, max_col), e
    coleta as mesclagens. Fórmulas são preservadas como texto "=...".
    """
    wb = load_workbook(BytesIO(conteudo), data_only=False)
    if sheet and sheet in wb.sheetnames:
        ws = wb[sheet]
    else:
        ws = wb.active

    max_row = ws.max_row or 0
    max_col = ws.max_column or 0

    cells: List[List[Optional[Dict[str, Any]]]] = []
    for r in range(1, max_row + 1):
        linha: List[Optional[Dict[str, Any]]] = []
        for c in range(1, max_col + 1):
            v = _cell_value(ws.cell(row=r, column=c).value)
            linha.append({"v": v} if v is not None else None)
        cells.append(linha)

    merges = [str(rng) for rng in ws.merged_cells.ranges]

    return {
        "sheet": ws.title,
        "n_rows": max_row,
        "n_cols": max_col,
        "cells": cells,
        "merges": merges,
        "meta": {},
    }


# ---------------------------------------------------------------------------
# grid -> XLSX
# ---------------------------------------------------------------------------
def grid_to_xlsx(grid: Dict[str, Any], model: Optional[Dict[str, Any]] = None) -> bytes:
    """Serializa o dict "grid" de volta para bytes .xlsx (round-trip).

    Escreve a matriz densa de células (`cells`), reaplica as mesclagens e o
    nome da aba. `model` é aceito por assinatura de contrato (para futura
    reconciliação de formato via estrutura) — quando ausente, o grid é escrito
    literalmente, o que garante o round-trip XLSX->grid->XLSX estável.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = str((grid or {}).get("sheet") or "Planilha")[:31]

    cells = (grid or {}).get("cells") or []
    for i, linha in enumerate(cells, start=1):
        if not linha:
            continue
        for j, cel in enumerate(linha, start=1):
            if cel is None:
                continue
            v = cel.get("v") if isinstance(cel, dict) else cel
            if v is None:
                continue
            ws.cell(row=i, column=j).value = v

    for rng in (grid or {}).get("merges") or []:
        try:
            ws.merge_cells(str(rng))
        except (ValueError, TypeError):
            continue

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


# ---------------------------------------------------------------------------
# BillingModel -> grid (SEM PII — só cabeçalhos/estrutura)
# ---------------------------------------------------------------------------
def _estrutura_headers_grid(estrutura: Dict[str, Any]) -> Dict[str, Any]:
    """Grid só com o cabeçalho de uma estrutura C1 (sem linhas de dado)."""
    try:
        data_row = int(estrutura.get("data_row"))
    except (TypeError, ValueError):
        header_rows = estrutura.get("header_rows") or [1]
        data_row = (max(header_rows) + 1) if header_rows else 2
    headers = estrutura.get("headers") or {}
    colunas = [c for c in (estrutura.get("colunas") or []) if isinstance(c, dict)]

    n_rows = max(1, data_row - 1)
    n_cols = 0
    for c in colunas:
        letra = c.get("letra")
        if letra:
            n_cols = max(n_cols, column_index_from_string(letra))
    for letra in headers:
        try:
            n_cols = max(n_cols, column_index_from_string(letra))
        except (ValueError, TypeError):
            continue
    n_cols = max(n_cols, 1)

    cells: List[List[Optional[Dict[str, Any]]]] = [
        [None] * n_cols for _ in range(n_rows)
    ]
    for letra, por_linha in headers.items():
        if not isinstance(por_linha, dict):
            continue
        try:
            ci = column_index_from_string(letra) - 1
        except (ValueError, TypeError):
            continue
        for linha, texto in por_linha.items():
            try:
                ri = int(linha) - 1
            except (TypeError, ValueError):
                continue
            if 0 <= ri < n_rows and 0 <= ci < n_cols:
                cells[ri][ci] = {"v": texto}

    merges = [
        str(m) for m in (estrutura.get("merges") or []) if isinstance(m, str)
    ]
    return {
        "sheet": str(estrutura.get("aba") or "Faturamento")[:31],
        "n_rows": n_rows,
        "n_cols": n_cols,
        "cells": cells,
        "merges": merges,
        "meta": {
            "header_rows": estrutura.get("header_rows") or [1],
            "data_row": data_row,
            "colunas": _colunas_meta_from_estrutura(colunas),
        },
    }


def _colunas_from_lista(colunas: List[str]) -> Dict[str, Any]:
    """Grid de cabeçalho para um modelo dirigido por lista de colunas (1 linha)."""
    colunas = [c for c in (colunas or []) if c]
    n_cols = max(len(colunas), 1)
    cells: List[List[Optional[Dict[str, Any]]]] = [
        [{"v": c} for c in colunas] or [None]
    ]
    meta_cols = [{"header": c, "tipo": "campo", "fonte": c} for c in colunas]
    return {
        "sheet": "Faturamento",
        "n_rows": 1,
        "n_cols": n_cols,
        "cells": cells,
        "merges": [],
        "meta": {"header_rows": [1], "data_row": 2, "colunas": meta_cols},
    }


def _colunas_meta_from_estrutura(colunas: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Extrai a lista de metadados de coluna (sem PII) de uma estrutura C1."""
    meta: List[Dict[str, Any]] = []
    for c in colunas:
        if not isinstance(c, dict):
            continue
        item = {"header": c.get("header"), "tipo": c.get("tipo")}
        if c.get("fonte"):
            item["fonte"] = c["fonte"]
        if c.get("template"):
            item["template"] = c["template"]
        if "valor" in c:
            item["valor"] = c["valor"]
        if c.get("letra"):
            item["letra"] = c["letra"]
        meta.append(item)
    return meta


def model_to_grid(model: Any) -> Dict[str, Any]:
    """Grid de CABEÇALHO do BillingModel (SEM PII — nenhuma linha de dado).

    Aceita um `BillingModel` (ORM) ou seu dict. Modelo com `estrutura` (upload)
    usa o cabeçalho C1; senão, monta uma linha com os nomes das colunas.
    """
    md = _as_model_dict(model)
    estrutura = md.get("estrutura")
    if isinstance(estrutura, dict) and estrutura.get("colunas"):
        return _estrutura_headers_grid(estrutura)
    colunas = md.get("colunas") or list(FEMSA_COLUMNS)
    return _colunas_from_lista(colunas)


# ---------------------------------------------------------------------------
# Snapshot (SEM PII — §5 do contrato)
# ---------------------------------------------------------------------------
def get_snapshot(model: Any) -> Dict[str, Any]:
    """GridSnapshot do modelo para o LLM — SÓ estrutura, SEM PII.

    Contém nomes de colunas, tipos, fórmulas e parâmetros do modelo; nunca CPF,
    nome, matrícula ou valores individuais reais. No máximo uma linha de exemplo
    SINTÉTICA (cabeçalhos), nunca dados de funcionário.
    """
    md = _as_model_dict(model)
    estrutura = md.get("estrutura")
    if isinstance(estrutura, dict) and estrutura.get("colunas"):
        colunas_meta = _colunas_meta_from_estrutura(
            [c for c in estrutura.get("colunas") or [] if isinstance(c, dict)]
        )
        colunas_nomes = [c.get("header") for c in colunas_meta]
    else:
        colunas_nomes = list(md.get("colunas") or FEMSA_COLUMNS)
        colunas_meta = [
            {"header": c, "tipo": "campo", "fonte": c} for c in colunas_nomes
        ]

    return {
        "model_id": md.get("id"),
        "nome": md.get("nome"),
        "colunas": colunas_nomes,
        "colunas_meta": colunas_meta,
        "campos_config": md.get("campos_config") or [],
        "params": {
            "encargos_pct": md.get("encargos_pct"),
            "taxa_adm_pct": md.get("taxa_adm_pct"),
            "imposto_pct": md.get("imposto_pct"),
            "salario_formula": md.get("salario_formula"),
        },
        "tem_estrutura": bool(estrutura),
        # Uma linha de exemplo SINTÉTICA (só rótulos), nunca PII:
        "sample_row": {c: None for c in colunas_nomes},
    }


# ---------------------------------------------------------------------------
# apply_ops (contrato §4)
# ---------------------------------------------------------------------------
def _as_model_dict(model: Any) -> Dict[str, Any]:
    """Aceita um BillingModel (ORM) ou um dict e devolve um dict de trabalho.

    Copia campos relevantes; para o ORM lê os atributos diretamente (evita
    depender da forma de `to_dict`, que omite estrutura/colunas por padrão).
    """
    if isinstance(model, dict):
        return copy.deepcopy(model)
    return {
        "id": getattr(model, "id", None),
        "nome": getattr(model, "nome", None),
        "colunas": copy.deepcopy(getattr(model, "colunas", None) or []),
        "estrutura": copy.deepcopy(getattr(model, "estrutura", None)),
        "campos_config": copy.deepcopy(getattr(model, "campos_config", None) or []),
        "encargos_pct": getattr(model, "encargos_pct", None),
        "taxa_adm_pct": getattr(model, "taxa_adm_pct", None),
        "imposto_pct": getattr(model, "imposto_pct", None),
        "salario_formula": getattr(model, "salario_formula", None),
    }


def _colunas_estrutura(md: Dict[str, Any]) -> Optional[List[Dict[str, Any]]]:
    """Lista de colunas da estrutura C1 do modelo (ou None se dirigido por lista)."""
    est = md.get("estrutura")
    if isinstance(est, dict) and isinstance(est.get("colunas"), list):
        return est["colunas"]
    return None


def _headers_atuais(md: Dict[str, Any]) -> List[str]:
    """Lista ordenada dos headers atuais (estrutura C1 ou lista de colunas)."""
    cols = _colunas_estrutura(md)
    if cols is not None:
        return [str(c.get("header") or "") for c in cols if isinstance(c, dict)]
    return list(md.get("colunas") or [])


def _find_header_idx(headers: List[str], header: str) -> int:
    for i, h in enumerate(headers):
        if h == header:
            return i
    return -1


def _next_letra(colunas: List[Dict[str, Any]]) -> str:
    """Próxima letra de coluna livre para uma estrutura C1."""
    usadas = set()
    for c in colunas:
        letra = c.get("letra")
        if letra:
            try:
                usadas.add(column_index_from_string(letra))
            except (ValueError, TypeError):
                continue
    i = 1
    while i in usadas:
        i += 1
    return get_column_letter(i)


def _op_add_column(md: Dict[str, Any], op: Dict[str, Any], res: ApplyResult) -> bool:
    header = str(op.get("header") or "").strip()
    if not header:
        res.errors.append("add_column: 'header' é obrigatório.")
        return False
    tipo = op.get("tipo") or "campo"
    if tipo not in ("campo", "formula", "constante", "vazio"):
        res.errors.append(f"add_column: tipo inválido '{tipo}'.")
        return False

    headers = _headers_atuais(md)
    if header in headers:
        res.errors.append(f"add_column: coluna '{header}' já existe.")
        return False

    novo: Dict[str, Any] = {"header": header, "tipo": tipo}
    if tipo == "formula":
        template = str(op.get("template") or "")
        if "{row}" not in template:
            res.errors.append("add_column: tipo 'formula' exige 'template' com '{row}'.")
            return False
        novo["template"] = template
    elif tipo == "constante":
        if "valor" not in op:
            res.errors.append("add_column: tipo 'constante' exige 'valor'.")
            return False
        novo["valor"] = op.get("valor")
    elif tipo == "campo":
        fonte = op.get("fonte") or header
        novo["fonte"] = fonte
        if fonte not in _KNOWN:
            res.warnings.append(f"add_column: fonte '{fonte}' não reconhecida — coluna sairá vazia.")

    after = op.get("after")

    cols = _colunas_estrutura(md)
    if cols is not None:
        novo["letra"] = _next_letra(cols)
        pos = len(cols)
        if after:
            idx = next((i for i, c in enumerate(cols)
                        if isinstance(c, dict) and c.get("header") == after), -1)
            if idx >= 0:
                pos = idx + 1
        cols.insert(pos, novo)
    else:
        # Modelo dirigido por lista: só faz sentido para colunas 'campo' conhecidas.
        colunas = md.setdefault("colunas", [])
        pos = len(colunas)
        if after and after in colunas:
            pos = colunas.index(after) + 1
        colunas.insert(pos, header)
    return True


def _op_remove_column(md: Dict[str, Any], op: Dict[str, Any], res: ApplyResult) -> bool:
    header = str(op.get("header") or "").strip()
    cols = _colunas_estrutura(md)
    if cols is not None:
        idx = next((i for i, c in enumerate(cols)
                    if isinstance(c, dict) and c.get("header") == header), -1)
        if idx < 0:
            res.errors.append(f"remove_column: coluna '{header}' não encontrada.")
            return False
        cols.pop(idx)
        return True
    colunas = md.get("colunas") or []
    if header not in colunas:
        res.errors.append(f"remove_column: coluna '{header}' não encontrada.")
        return False
    colunas.remove(header)
    md["colunas"] = colunas
    return True


def _op_rename_column(md: Dict[str, Any], op: Dict[str, Any], res: ApplyResult) -> bool:
    header = str(op.get("header") or "").strip()
    novo = str(op.get("novo_header") or "").strip()
    if not novo:
        res.errors.append("rename_column: 'novo_header' é obrigatório.")
        return False
    headers = _headers_atuais(md)
    if novo in headers and novo != header:
        res.errors.append(f"rename_column: já existe uma coluna '{novo}'.")
        return False
    cols = _colunas_estrutura(md)
    if cols is not None:
        c = next((c for c in cols if isinstance(c, dict) and c.get("header") == header), None)
        if c is None:
            res.errors.append(f"rename_column: coluna '{header}' não encontrada.")
            return False
        c["header"] = novo
        return True
    colunas = md.get("colunas") or []
    if header not in colunas:
        res.errors.append(f"rename_column: coluna '{header}' não encontrada.")
        return False
    colunas[colunas.index(header)] = novo
    md["colunas"] = colunas
    return True


def _op_move_column(md: Dict[str, Any], op: Dict[str, Any], res: ApplyResult) -> bool:
    header = str(op.get("header") or "").strip()
    before = op.get("before")
    after = op.get("after")
    cols = _colunas_estrutura(md)
    if cols is not None:
        idx = next((i for i, c in enumerate(cols)
                    if isinstance(c, dict) and c.get("header") == header), -1)
        if idx < 0:
            res.errors.append(f"move_column: coluna '{header}' não encontrada.")
            return False
        item = cols.pop(idx)
        pos = _move_target_pos([c.get("header") for c in cols], before, after, res)
        if pos is None:
            cols.insert(idx, item)  # desfaz
            return False
        cols.insert(pos, item)
        return True
    colunas = md.get("colunas") or []
    if header not in colunas:
        res.errors.append(f"move_column: coluna '{header}' não encontrada.")
        return False
    colunas.remove(header)
    pos = _move_target_pos(colunas, before, after, res)
    if pos is None:
        colunas.insert(0, header)
        return False
    colunas.insert(pos, header)
    md["colunas"] = colunas
    return True


def _move_target_pos(headers: List[str], before: Any, after: Any,
                     res: ApplyResult) -> Optional[int]:
    """Índice de inserção a partir de before/after; None (com erro) se o alvo some."""
    if before:
        idx = _find_header_idx(headers, before)
        if idx < 0:
            res.errors.append(f"move_column: alvo 'before={before}' não encontrado.")
            return None
        return idx
    if after:
        idx = _find_header_idx(headers, after)
        if idx < 0:
            res.errors.append(f"move_column: alvo 'after={after}' não encontrado.")
            return None
        return idx + 1
    return len(headers)  # sem alvo -> vai para o fim


def _op_set_field_config(md: Dict[str, Any], op: Dict[str, Any], res: ApplyResult) -> bool:
    campo = str(op.get("campo") or "").strip()
    if not campo:
        res.errors.append("set_field_config: 'campo' é obrigatório.")
        return False
    if campo not in _headers_atuais(md) and campo not in _KNOWN:
        res.warnings.append(f"set_field_config: campo '{campo}' não está entre as colunas do modelo.")
    codigo = op.get("codigo")
    if codigo is not None and str(codigo).strip():
        for parte in str(codigo).split(","):
            if not parte.strip().isdigit():
                res.errors.append(f"set_field_config: código inválido '{parte}' (use números separados por vírgula).")
                return False
    formula = op.get("formula")
    if formula is not None and str(formula).strip():
        from app.services.formula_salario import validar_formula
        erro = validar_formula(str(formula).strip())
        if erro:
            res.errors.append(f"set_field_config ({campo}): {erro}")
            return False

    cfg = md.setdefault("campos_config", [])
    existente = next((c for c in cfg if isinstance(c, dict) and c.get("campo") == campo), None)
    linha = {
        "campo": campo,
        "codigo": (",".join(p.strip() for p in str(codigo).split(",") if p.strip())
                   if codigo not in (None, "") else None),
        "codigo_nome": (str(op.get("codigo_nome")).strip() or None) if op.get("codigo_nome") else None,
        "formula": (str(formula).strip() or None) if formula else None,
    }
    if existente is not None:
        existente.update(linha)
    else:
        cfg.append(linha)
    return True


def _op_set_model_param(md: Dict[str, Any], op: Dict[str, Any], res: ApplyResult) -> bool:
    param = str(op.get("param") or "").strip()
    if param not in _PARAM_TO_FIELD:
        res.errors.append(f"set_model_param: parâmetro inválido '{param}'.")
        return False
    valor = op.get("valor")
    if param in _PCT_PARAMS:
        try:
            fval = float(valor)
        except (TypeError, ValueError):
            res.errors.append(f"set_model_param: '{param}' deve ser numérico (0–100).")
            return False
        if not (0 <= fval <= 100):
            res.errors.append(f"set_model_param: '{param}' deve estar entre 0 e 100.")
            return False
        md[_PARAM_TO_FIELD[param]] = fval
    else:  # salario_formula
        if valor in (None, ""):
            md[_PARAM_TO_FIELD[param]] = None
        else:
            from app.services.formula_salario import validar_formula
            erro = validar_formula(str(valor).strip())
            if erro:
                res.errors.append(f"set_model_param: fórmula do salário inválida: {erro}")
                return False
            md[_PARAM_TO_FIELD[param]] = str(valor).strip()
    return True


def _op_add_item(md: Dict[str, Any], op: Dict[str, Any], res: ApplyResult) -> bool:
    """add_item (contrato §4/§5): opera sobre estrutura/dados sintéticos.

    Nesta camada de MODELO não injetamos PII nem persistimos linhas de dado
    reais — reconhecemos os targets válidos e devolvemos um warning
    informativo. Alvos desconhecidos são erro.
    """
    target = str(op.get("target") or "").strip()
    if target not in ("catalogo", "linha"):
        res.errors.append(f"add_item: target inválido '{target}' (use 'catalogo' ou 'linha').")
        return False
    res.warnings.append(
        f"add_item(target={target}) registrado como operação sintética — "
        "não injeta PII no modelo (§5 do contrato)."
    )
    return True


_OP_HANDLERS = {
    "add_column": _op_add_column,
    "remove_column": _op_remove_column,
    "rename_column": _op_rename_column,
    "move_column": _op_move_column,
    "set_field_config": _op_set_field_config,
    "set_model_param": _op_set_model_param,
    "add_item": _op_add_item,
}


def apply_ops(model: Any, ops: List[Dict[str, Any]]) -> ApplyResult:
    """Valida e aplica uma lista de GridOp sobre o modelo (contrato §4/§6).

    Aplica sobre uma CÓPIA do modelo — nada é mutado no objeto de entrada. Se
    QUALQUER op falhar na validação, NADA é aplicado (all-or-nothing): o
    resultado volta com `ok=False`, os `errors` e o modelo ORIGINAL intacto,
    para nunca corromper o modelo (edge case da spec).
    """
    md = _as_model_dict(model)
    trabalho = copy.deepcopy(md)
    res = ApplyResult(model=md)

    if not isinstance(ops, list):
        res.ok = False
        res.errors.append("Payload inválido: 'ops' deve ser uma lista.")
        return res

    applied = 0
    for i, op in enumerate(ops, start=1):
        if not isinstance(op, dict):
            res.errors.append(f"Op #{i}: formato inválido (esperado objeto).")
            continue
        nome = op.get("op")
        handler = _OP_HANDLERS.get(nome)
        if handler is None:
            res.errors.append(f"Op #{i}: operação desconhecida '{nome}'.")
            continue
        antes_erros = len(res.errors)
        ok = handler(trabalho, op, res)
        if ok and len(res.errors) == antes_erros:
            applied += 1

    if res.errors:
        # All-or-nothing: não aplica parcialmente sem relatar (contrato §4).
        res.ok = False
        res.applied = 0
        res.model = md  # modelo original, intacto
        return res

    # Se o modelo tem estrutura, revalida (recusa se ficou inconsistente).
    est = trabalho.get("estrutura")
    if isinstance(est, dict) and est.get("colunas"):
        problemas = validate_estrutura(est)
        # Só bloqueia problemas estruturais graves; fonte desconhecida já virou
        # warning na add_column. Filtra os avisos de fonte que aceitamos.
        graves = [p for p in problemas if "fonte desconhecida" not in p.lower()]
        if graves:
            res.ok = False
            res.applied = 0
            res.errors.extend(graves)
            res.model = md
            return res
        # Mantém colunas (lista canônica) coerente com a estrutura.
        trabalho["colunas"] = derive_colunas(est)

    res.ok = True
    res.applied = applied
    res.model = trabalho
    return res


# ---------------------------------------------------------------------------
# Igualdade tolerante de grids (para testes de round-trip)
# ---------------------------------------------------------------------------
def grids_equivalentes(a: Dict[str, Any], b: Dict[str, Any]) -> Tuple[bool, str]:
    """Compara dois grids de forma tolerante (útil no round-trip e preview×export).

    Ignora diferenças irrelevantes: 5 vs 5.0, espaços em texto, e sobra de
    linhas/colunas totalmente vazias no fim. Retorna (igual, motivo).
    """
    ca = _densa_norm((a or {}).get("cells") or [])
    cb = _densa_norm((b or {}).get("cells") or [])
    if ca != cb:
        # localiza a primeira divergência para a mensagem
        for r in range(max(len(ca), len(cb))):
            la = ca[r] if r < len(ca) else []
            lb = cb[r] if r < len(cb) else []
            for c in range(max(len(la), len(lb))):
                va = la[c] if c < len(la) else None
                vb = lb[c] if c < len(lb) else None
                if va != vb:
                    return False, f"célula ({r + 1},{c + 1}): {va!r} != {vb!r}"
        return False, "matrizes diferem"
    if set((a or {}).get("merges") or []) != set((b or {}).get("merges") or []):
        return False, "merges diferentes"
    return True, ""


def _densa_norm(cells: List[List[Optional[Dict[str, Any]]]]) -> List[List[Any]]:
    """Normaliza a matriz para comparação: valores normalizados + trim de vazios."""
    norm: List[List[Any]] = []
    for linha in cells:
        nl = [_norm_num((cel or {}).get("v") if isinstance(cel, dict) else cel)
              for cel in (linha or [])]
        while nl and (nl[-1] is None or nl[-1] == ""):
            nl.pop()
        norm.append(nl)
    while norm and not norm[-1]:
        norm.pop()
    return norm
